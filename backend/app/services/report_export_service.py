"""
PDF and Excel Report Export Service for OBE Framework.

Generates professional CO-PO-PSO attainment reports using:
  • reportlab for PDF
  • openpyxl  for Excel
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any, Dict, List, Optional


# ─────────────────────────────────────────────────────────────────────────────
# PDF generation
# ─────────────────────────────────────────────────────────────────────────────

def _level_color_pdf(level: str):
    """Return a reportlab Color for an attainment level."""
    from reportlab.lib import colors
    mapping = {
        "Level 3": colors.HexColor("#27AE60"),   # green
        "Level 2": colors.HexColor("#F39C12"),   # amber
        "Level 1": colors.HexColor("#E74C3C"),   # red
    }
    return mapping.get(level, colors.lightgrey)


def generate_pdf_report(course_data: Dict[str, Any]) -> bytes:
    """Render a complete OBE attainment report as PDF bytes."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        rightMargin=0.6 * inch, leftMargin=0.6 * inch,
        topMargin=0.6 * inch, bottomMargin=0.6 * inch,
    )

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=16, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, spaceAfter=4)
    normal = styles["Normal"]

    elements: list = []

    # ── Title ────────────────────────────────────────────────────────────────
    elements.append(Paragraph("CO-PO-PSO Attainment Report", h1))
    elements.append(Paragraph(
        f"Course: <b>{course_data.get('course_code', '')} – "
        f"{course_data.get('course_name', '')}</b>", normal
    ))
    elements.append(Paragraph(
        f"Generated: {datetime.utcnow().strftime('%d %b %Y, %H:%M UTC')}", normal
    ))
    elements.append(Spacer(1, 0.2 * inch))
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.grey))
    elements.append(Spacer(1, 0.15 * inch))

    # ── CO Attainment ────────────────────────────────────────────────────────
    cos = course_data.get("co_attainments", [])
    if cos:
        elements.append(Paragraph("Course Outcome (CO) Attainment", h2))
        header = [
            _bold_cell("CO Code"), _bold_cell("CO Statement"),
            _bold_cell("Attainment %"), _bold_cell("Level"),
            _bold_cell("Students"),
        ]
        rows = [header]
        for co in cos:
            pct = co.get("attainment_percentage", 0)
            lvl = co.get("attainment_level", "Level 1")
            stmt = co.get("co_statement", co.get("statement", ""))
            rows.append([
                co.get("co_code", ""),
                stmt[:70] + ("…" if len(stmt) > 70 else ""),
                f"{pct:.1f}%",
                lvl,
                str(co.get("students", co.get("total_students", "–"))),
            ])
        tbl = Table(rows, colWidths=[0.8 * inch, 3.4 * inch, 1.0 * inch, 0.9 * inch, 0.7 * inch])
        _apply_base_style(tbl, colors.HexColor("#1A5276"), len(rows))
        # Colour the attainment level cells
        for i, co in enumerate(cos, start=1):
            lvl = co.get("attainment_level", "Level 1")
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (3, i), (3, i), _level_color_pdf(lvl)),
                ("TEXTCOLOR", (3, i), (3, i), colors.white),
            ]))
        elements += [tbl, Spacer(1, 0.2 * inch)]

        # Summary line
        avg = course_data.get("average_co_attainment",
                              sum(c.get("attainment_percentage", 0) for c in cos) / max(len(cos), 1))
        elements.append(Paragraph(
            f"<b>Average CO Attainment: {avg:.1f}%  •  "
            f"Overall Level: {_level_text(avg)}</b>", normal
        ))
        elements.append(Spacer(1, 0.15 * inch))

    # ── PO Attainment ────────────────────────────────────────────────────────
    pos = course_data.get("po_attainments", [])
    if pos:
        elements.append(Paragraph("Program Outcome (PO) Attainment", h2))
        header = [
            _bold_cell("PO Code"), _bold_cell("PO Statement"),
            _bold_cell("Attainment %"), _bold_cell("Level"),
            _bold_cell("Mapped COs"),
        ]
        rows = [header]
        for po in pos:
            pct = po.get("attainment_percentage", 0)
            lvl = po.get("attainment_level", "Level 1")
            stmt = po.get("po_statement", po.get("statement", ""))
            rows.append([
                po.get("po_code", ""),
                stmt[:70] + ("…" if len(stmt) > 70 else ""),
                f"{pct:.1f}%",
                lvl,
                str(po.get("mapped_cos", "–")),
            ])
        tbl = Table(rows, colWidths=[0.8 * inch, 3.4 * inch, 1.0 * inch, 0.9 * inch, 0.7 * inch])
        _apply_base_style(tbl, colors.HexColor("#1E8449"), len(rows))
        for i, po in enumerate(pos, start=1):
            lvl = po.get("attainment_level", "Level 1")
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (3, i), (3, i), _level_color_pdf(lvl)),
                ("TEXTCOLOR", (3, i), (3, i), colors.white),
            ]))
        elements += [tbl, Spacer(1, 0.2 * inch)]

    # ── PSO Attainment ───────────────────────────────────────────────────────
    psos = course_data.get("pso_attainments", [])
    if psos:
        elements.append(Paragraph("Program Specific Outcome (PSO) Attainment", h2))
        header = [
            _bold_cell("PSO Code"), _bold_cell("Attainment %"),
            _bold_cell("Level"), _bold_cell("Mapped COs"),
        ]
        rows = [header]
        for pso in psos:
            pct = pso.get("attainment_percentage", 0)
            lvl = pso.get("attainment_level", "Level 1")
            rows.append([
                pso.get("pso_code", pso.get("code", "")),
                f"{pct:.1f}%", lvl,
                str(pso.get("mapped_cos", "–")),
            ])
        tbl = Table(rows, colWidths=[1.0 * inch, 1.5 * inch, 1.2 * inch, 1.5 * inch])
        _apply_base_style(tbl, colors.HexColor("#6E2FA1"), len(rows))
        elements += [tbl, Spacer(1, 0.2 * inch)]

    # ── CO-PO Mapping Matrix ─────────────────────────────────────────────────
    matrix = course_data.get("co_po_matrix")
    if matrix:
        elements.append(Paragraph("CO-PO Mapping Matrix (Correlation Level)", h2))
        cos_list = matrix.get("cos", [])
        pos_list = matrix.get("pos", [])
        data = matrix.get("data", {})
        header = [""] + pos_list
        rows = [header]
        for co_code in cos_list:
            row = [co_code] + [
                str(data.get(co_code, {}).get(po_code, "-"))
                for po_code in pos_list
            ]
            rows.append(row)
        n_cols = len(header)
        col_w = min(0.7 * inch, 5.5 * inch / max(n_cols, 1))
        tbl = Table(rows, colWidths=[col_w] * n_cols)
        _apply_base_style(tbl, colors.HexColor("#2E4057"), len(rows))
        elements += [tbl, Spacer(1, 0.2 * inch)]

    doc.build(elements)
    return buffer.getvalue()


def _bold_cell(text: str) -> str:
    return f"<b>{text}</b>"


def _apply_base_style(tbl: Any, header_color: Any, n_rows: int) -> None:
    from reportlab.lib import colors
    from reportlab.platypus import TableStyle
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), header_color),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("ALIGN", (1, 1), (1, -1), "LEFT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#EAF0FA"), colors.white]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))


def _level_text(pct: float) -> str:
    if pct >= 70:
        return "Level 3 (Attained)"
    elif pct >= 60:
        return "Level 2 (Partially Attained)"
    return "Level 1 (Not Attained)"


# ─────────────────────────────────────────────────────────────────────────────
# Excel generation
# ─────────────────────────────────────────────────────────────────────────────

def _xlsx_header_style(ws: Any, row: int, headers: List[str], fill_hex: str) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = max(
            ws.column_dimensions[get_column_letter(col)].width or 0,
            len(h) + 4,
        )


def _xlsx_level_fill(level: str) -> Any:
    from openpyxl.styles import PatternFill
    mapping = {
        "Level 3": PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid"),
        "Level 2": PatternFill(start_color="F39C12", end_color="F39C12", fill_type="solid"),
        "Level 1": PatternFill(start_color="E74C3C", end_color="E74C3C", fill_type="solid"),
    }
    return mapping.get(level)


def generate_excel_report(course_data: Dict[str, Any]) -> bytes:
    """Render a complete OBE attainment report as xlsx bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum["A1"] = "CO-PO-PSO Attainment Report"
    ws_sum["A1"].font = Font(bold=True, size=14, color="1A5276")
    ws_sum["A2"] = f"Course: {course_data.get('course_code', '')} – {course_data.get('course_name', '')}"
    ws_sum["A3"] = f"Generated: {datetime.utcnow().strftime('%d %b %Y %H:%M UTC')}"
    ws_sum["A4"] = f"Avg CO Attainment: {course_data.get('average_co_attainment', 0):.1f}%"
    ws_sum["A5"] = f"Avg PO Attainment: {course_data.get('average_po_attainment', 0):.1f}%"
    ws_sum["A6"] = f"Overall Level: {course_data.get('overall_level', '–')}"

    # ── Sheet 2: CO Attainment ────────────────────────────────────────────────
    ws1 = wb.create_sheet("CO Attainment")
    co_headers = ["CO Code", "CO Statement", "Attainment %", "Level",
                  "Total Marks", "Marks Obtained", "Students"]
    _xlsx_header_style(ws1, 1, co_headers, "1A5276")

    for row_idx, co in enumerate(course_data.get("co_attainments", []), 2):
        lvl = co.get("attainment_level", "Level 1")
        stmt = co.get("co_statement", co.get("statement", ""))
        ws1.cell(row=row_idx, column=1, value=co.get("co_code", ""))
        ws1.cell(row=row_idx, column=2, value=stmt)
        pct_cell = ws1.cell(row=row_idx, column=3, value=round(co.get("attainment_percentage", 0), 2))
        pct_cell.number_format = "0.00"
        lvl_cell = ws1.cell(row=row_idx, column=4, value=lvl)
        fill = _xlsx_level_fill(lvl)
        if fill:
            lvl_cell.fill = fill
            lvl_cell.font = Font(bold=True, color="FFFFFF")
        ws1.cell(row=row_idx, column=5, value=co.get("total_marks", 0))
        ws1.cell(row=row_idx, column=6, value=round(co.get("total_obtained", co.get("marks_obtained", 0)), 2))
        ws1.cell(row=row_idx, column=7, value=co.get("students", co.get("total_students", 0)))

    ws1.column_dimensions["B"].width = 50

    # ── Sheet 3: PO Attainment ────────────────────────────────────────────────
    ws2 = wb.create_sheet("PO Attainment")
    po_headers = ["PO Code", "PO Statement", "Attainment %", "Level", "Mapped COs"]
    _xlsx_header_style(ws2, 1, po_headers, "1E8449")

    for row_idx, po in enumerate(course_data.get("po_attainments", []), 2):
        lvl = po.get("attainment_level", "Level 1")
        stmt = po.get("po_statement", po.get("statement", ""))
        ws2.cell(row=row_idx, column=1, value=po.get("po_code", ""))
        ws2.cell(row=row_idx, column=2, value=stmt)
        pct_cell = ws2.cell(row=row_idx, column=3, value=round(po.get("attainment_percentage", 0), 2))
        pct_cell.number_format = "0.00"
        lvl_cell = ws2.cell(row=row_idx, column=4, value=lvl)
        fill = _xlsx_level_fill(lvl)
        if fill:
            lvl_cell.fill = fill
            lvl_cell.font = Font(bold=True, color="FFFFFF")
        ws2.cell(row=row_idx, column=5, value=po.get("mapped_cos", 0))

    ws2.column_dimensions["B"].width = 50

    # ── Sheet 4: PSO Attainment ───────────────────────────────────────────────
    psos = course_data.get("pso_attainments", [])
    if psos:
        ws3 = wb.create_sheet("PSO Attainment")
        pso_headers = ["PSO Code", "PSO Statement", "Attainment %", "Level", "Mapped COs"]
        _xlsx_header_style(ws3, 1, pso_headers, "6E2FA1")
        for row_idx, pso in enumerate(psos, 2):
            lvl = pso.get("attainment_level", "Level 1")
            ws3.cell(row=row_idx, column=1, value=pso.get("pso_code", pso.get("code", "")))
            ws3.cell(row=row_idx, column=2, value=pso.get("pso_statement", pso.get("statement", "")))
            ws3.cell(row=row_idx, column=3, value=round(pso.get("attainment_percentage", 0), 2)).number_format = "0.00"
            lvl_cell = ws3.cell(row=row_idx, column=4, value=lvl)
            fill = _xlsx_level_fill(lvl)
            if fill:
                lvl_cell.fill = fill
                lvl_cell.font = Font(bold=True, color="FFFFFF")
            ws3.cell(row=row_idx, column=5, value=pso.get("mapped_cos", 0))

    # ── Sheet 5: CO-PO Matrix ─────────────────────────────────────────────────
    matrix = course_data.get("co_po_matrix")
    if matrix:
        ws_m = wb.create_sheet("CO-PO Matrix")
        cos_list = matrix.get("cos", [])
        pos_list = matrix.get("pos", [])
        data = matrix.get("data", {})

        # Header row: PO codes
        ws_m.cell(row=1, column=1, value="CO\\PO").font = Font(bold=True)
        for j, po_code in enumerate(pos_list, 2):
            cell = ws_m.cell(row=1, column=j, value=po_code)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        for i, co_code in enumerate(cos_list, 2):
            ws_m.cell(row=i, column=1, value=co_code).font = Font(bold=True)
            for j, po_code in enumerate(pos_list, 2):
                val = data.get(co_code, {}).get(po_code, "-")
                cell = ws_m.cell(row=i, column=j, value=val)
                cell.alignment = Alignment(horizontal="center")
                # Colour by mapping level
                if val == 3:
                    cell.fill = PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid")
                    cell.font = Font(bold=True, color="FFFFFF")
                elif val == 2:
                    cell.fill = PatternFill(start_color="F39C12", end_color="F39C12", fill_type="solid")
                    cell.font = Font(bold=True, color="FFFFFF")
                elif val == 1:
                    cell.fill = PatternFill(start_color="AED6F1", end_color="AED6F1", fill_type="solid")

    # ── Sheet 6: Student Performance ─────────────────────────────────────────
    student_perf = course_data.get("student_performance", [])
    if student_perf:
        ws_sp = wb.create_sheet("Student Performance")
        # student_performance: [{student_id, total_marks, percentage, co_marks: {co_code: marks}}]
        co_codes = list({
            co_code
            for s in student_perf
            for co_code in s.get("co_marks", {}).keys()
        })
        sp_headers = ["Student ID", "Total Marks", "Percentage (%)"] + co_codes
        _xlsx_header_style(ws_sp, 1, sp_headers, "2C3E50")
        for row_idx, s in enumerate(student_perf, 2):
            ws_sp.cell(row=row_idx, column=1, value=s.get("student_id", ""))
            ws_sp.cell(row=row_idx, column=2, value=round(s.get("total_marks", 0), 2))
            pct = s.get("percentage", 0)
            pct_cell = ws_sp.cell(row=row_idx, column=3, value=round(pct, 2))
            pct_cell.number_format = "0.00"
            # Colour row by performance
            if pct >= 70:
                row_fill = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")
            elif pct >= 60:
                row_fill = PatternFill(start_color="FDEBD0", end_color="FDEBD0", fill_type="solid")
            else:
                row_fill = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
            for col in range(1, 4):
                ws_sp.cell(row=row_idx, column=col).fill = row_fill
            for c_idx, co_code in enumerate(co_codes, 4):
                ws_sp.cell(row=row_idx, column=c_idx, value=s.get("co_marks", {}).get(co_code, "–"))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generate_nba_excel_report(course_data: Dict[str, Any]) -> bytes:
    """
    Render NBA-style attainment workbook with target-vs-achieved sections.

    This format keeps tabular evidence that is typically requested for OBE audits.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "NBA Summary"

    ws["A1"] = "NBA OBE Attainment Report"
    ws["A1"].font = Font(bold=True, size=15, color="1F4E78")
    ws["A2"] = f"Course Code: {course_data.get('course_code', '')}"
    ws["A3"] = f"Course Name: {course_data.get('course_name', '')}"
    ws["A4"] = f"Generated On: {datetime.utcnow().strftime('%d %b %Y %H:%M UTC')}"
    ws["A6"] = "NBA thresholds: Level 3 >= 70%, Level 2 >= 60%, Level 1 < 60%"

    headers = [
        "Outcome Type",
        "Outcome Code",
        "Target %",
        "Achieved %",
        "Attainment Level",
        "Status",
    ]
    for idx, header in enumerate(headers, 1):
        cell = ws.cell(row=8, column=idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    target_pct = 70.0
    row = 9

    def _status_for_pct(pct: float) -> str:
        return "Target Met" if pct >= target_pct else "Action Required"

    for co in course_data.get("co_attainments", []):
        pct = float(co.get("attainment_percentage", 0.0))
        ws.cell(row=row, column=1, value="CO")
        ws.cell(row=row, column=2, value=co.get("co_code", ""))
        ws.cell(row=row, column=3, value=target_pct)
        ws.cell(row=row, column=4, value=round(pct, 2))
        ws.cell(row=row, column=5, value=co.get("attainment_level", "Level 1"))
        ws.cell(row=row, column=6, value=_status_for_pct(pct))
        row += 1

    for po in course_data.get("po_attainments", []):
        pct = float(po.get("attainment_percentage", 0.0))
        ws.cell(row=row, column=1, value="PO")
        ws.cell(row=row, column=2, value=po.get("po_code", ""))
        ws.cell(row=row, column=3, value=target_pct)
        ws.cell(row=row, column=4, value=round(pct, 2))
        ws.cell(row=row, column=5, value=po.get("attainment_level", "Level 1"))
        ws.cell(row=row, column=6, value=_status_for_pct(pct))
        row += 1

    for pso in course_data.get("pso_attainments", []):
        pct = float(pso.get("attainment_percentage", 0.0))
        ws.cell(row=row, column=1, value="PSO")
        ws.cell(row=row, column=2, value=pso.get("pso_code", pso.get("code", "")))
        ws.cell(row=row, column=3, value=target_pct)
        ws.cell(row=row, column=4, value=round(pct, 2))
        ws.cell(row=row, column=5, value=pso.get("attainment_level", "Level 1"))
        ws.cell(row=row, column=6, value=_status_for_pct(pct))
        row += 1

    # Add improvement notes section for audit readiness.
    ws.cell(row=row + 2, column=1, value="Continuous Improvement Notes").font = Font(bold=True, color="1F4E78")
    ws.cell(row=row + 3, column=1, value="1. Identify outcomes below target and map interventions.")
    ws.cell(row=row + 4, column=1, value="2. Re-evaluate after corrective actions in next cycle.")

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 18

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
